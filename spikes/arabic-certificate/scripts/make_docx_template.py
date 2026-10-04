"""ينشئ قالب DOCX عربياً (RTL) بمتغيرات {name} يملؤها docxtemplater. يُشغَّل مرة واحدة."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor

FONT = "Amiri"
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(16)
sec.top_margin, sec.bottom_margin = Mm(18), Mm(16)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(14)
rpr = st.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts")
    rpr.append(rf)
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)
szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "28"); rpr.append(szcs)


def rtl_par(text="", bold=False, size=None, align=None, color=None):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))
    if align is not None:
        p.alignment = align  # بلا jc = بداية السطر (يمين) في الفقرات RTL؛ jc=right ينقلب يساراً
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        rr = r._r.get_or_add_rPr()
        rr.append(OxmlElement("w:rtl"))
        if bold:
            rr.append(OxmlElement("w:bCs"))
    return p


rtl_par("{company_name}", True, 17, color="1D3557")
rtl_par("{company_address}", size=11)
rtl_par("الرقم: {doc_no}      التاريخ: {issue_date} م  |  {issue_date_hijri}", size=12)
rtl_par("{doc_title}", True, 20, WD_ALIGN_PARAGRAPH.CENTER)
rtl_par("{addressee}", True)
rtl_par("تحية طيبة وبعد،")
rtl_par("نشهد نحن شركة {company_name} بأن السيد/ة {employee_name}، الحامل/ة للرقم الوطني {national_id}، "
        "يعمل/تعمل لدينا بوظيفة {job_title} في {department} منذ تاريخ {hire_date}، "
        "وما زال/زالت على رأس عمله/عملها حتى تاريخه.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

rows = [("الرقم الوظيفي", "{employee_no}"), ("المسمى الوظيفي", "{job_title}"), ("القسم", "{department}"),
        ("تاريخ التعيين", "{hire_date}"), ("الراتب الشهري الإجمالي", "{salary} {currency}")]
tbl = doc.add_table(rows=len(rows), cols=2)
tbl.style = "Table Grid"
tblpr = tbl._tbl.tblPr
tblpr.append(OxmlElement("w:bidiVisual"))
for i, (k, v) in enumerate(rows):
    for j, txt in enumerate((k, v)):
        c = tbl.cell(i, j)
        c.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        run = c.paragraphs[0].add_run(txt)
        run.bold = j == 0
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))

rtl_par("وقد أُعطيت هذه الشهادة بناءً على طلبه/طلبها، {purpose}.")
rtl_par("وتفضلوا بقبول فائق الاحترام والتقدير،،،")
rtl_par("{signatory_name}", True)
rtl_par("{signatory_title}")
rtl_par("التوقيع والختم: ____________________")
rtl_par("للتواصل: {email}   |   رقم الوثيقة {doc_no}", size=10)

doc.save("templates/certificate.ar.docx")
print("saved templates/certificate.ar.docx")
